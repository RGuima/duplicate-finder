"""Attribute extraction for near-duplicate detection.

For photos/videos, "very similar" means same capture attributes (date
taken, camera make/model) despite a different filename or a re-saved copy
with a different hash. For documents it means matching author/creation
metadata. Each extractor degrades gracefully (returns no signature) when
the relevant library isn't installed or the file has no usable metadata --
those files still get exact-duplicate detection via hashing, they just
won't be grouped as "similar".
"""
from __future__ import annotations

import json
import os


def _extract_image(path: str) -> dict:
    from PIL import Image

    meta: dict = {}
    with Image.open(path) as img:
        meta["width"], meta["height"] = img.size
        exif = img.getexif()
        if exif:
            make = exif.get(271)
            model = exif.get(272)
            date_taken = exif.get(306)
            try:
                exif_ifd = exif.get_ifd(0x8769)
            except Exception:
                exif_ifd = {}
            date_taken = exif_ifd.get(36867) or exif_ifd.get(36868) or date_taken
            if make:
                meta["make"] = str(make).strip()
            if model:
                meta["model"] = str(model).strip()
            if date_taken:
                meta["date_taken"] = str(date_taken).strip()
    return meta


def _extract_video(path: str) -> dict:
    meta: dict = {}
    try:
        from hachoir.parser import createParser
        from hachoir.metadata import extractMetadata
    except ImportError:
        return meta
    try:
        parser = createParser(path)
        if not parser:
            return meta
        with parser:
            md = extractMetadata(parser)
        if md:
            if md.has("creation_date"):
                meta["date_taken"] = str(md.get("creation_date"))
            if md.has("duration"):
                meta["duration"] = str(md.get("duration"))
            if md.has("width"):
                meta["width"] = md.get("width")
            if md.has("height"):
                meta["height"] = md.get("height")
    except Exception:
        pass
    return meta


def _extract_doc(path: str, ext: str) -> dict:
    meta: dict = {}
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            r = PdfReader(path)
            info = r.metadata
            if info:
                if info.author:
                    meta["author"] = info.author.strip()
                if info.creation_date:
                    meta["created"] = str(info.creation_date)
            meta["pages"] = len(r.pages)
        elif ext == ".docx":
            import docx

            cp = docx.Document(path).core_properties
            if cp.author:
                meta["author"] = cp.author.strip()
            if cp.created:
                meta["created"] = str(cp.created)
            if cp.title:
                meta["title"] = cp.title
        elif ext in (".xlsx",):
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True)
            props = wb.properties
            if props.creator:
                meta["author"] = props.creator.strip()
            if props.created:
                meta["created"] = str(props.created)
            wb.close()
        elif ext == ".pptx":
            from pptx import Presentation

            cp = Presentation(path).core_properties
            if cp.author:
                meta["author"] = cp.author.strip()
            if cp.created:
                meta["created"] = str(cp.created)
    except Exception:
        pass
    return meta


def _signature(category: str, meta: dict) -> str | None:
    if category == "image":
        date_taken = meta.get("date_taken")
        make = meta.get("make")
        model = meta.get("model")
        if date_taken and (make or model):
            return f"img|{date_taken}|{make or ''}|{model or ''}"
        if date_taken:
            return f"img|{date_taken}|{meta.get('width')}x{meta.get('height')}"
        return None
    if category == "video":
        date_taken = meta.get("date_taken")
        duration = meta.get("duration")
        if date_taken and duration:
            return f"vid|{date_taken}|{duration}"
        if date_taken:
            return f"vid|{date_taken}"
        return None
    if category == "doc":
        author = meta.get("author")
        created = meta.get("created")
        if author and created:
            return f"doc|{author}|{created}"
        if created:
            return f"doc|{created}"
        return None
    return None


def extract(path: str, category: str) -> tuple[str, str | None, str | None, str | None]:
    """Returns (path, meta_json, signature, error)."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if category == "image":
            meta = _extract_image(path)
        elif category == "video":
            meta = _extract_video(path)
        elif category == "doc":
            meta = _extract_doc(path, ext)
        else:
            meta = {}
        sig = _signature(category, meta)
        return path, json.dumps(meta), sig, None
    except Exception as e:
        return path, json.dumps({}), None, str(e)
